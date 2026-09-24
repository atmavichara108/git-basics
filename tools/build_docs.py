#!/usr/bin/env python3
"""Сборка документов для отправки компании из исходников репозитория.

Исходники (редактировать их):
    02_questions/questionnaire.md      — вопросы к компании
    02_questions/documents_request.md  — перечень запрашиваемых документов
    03_roles/roles.csv                 — справочник специальностей

Результат (не редактировать руками — перезапишется):
    02_questions/voprosy_kompanii.docx — опросник для компании (Word, с колонкой для ответов)
    02_questions/zayavka_na_podbor.xlsx — таблица-заявка по специальностям (Excel)
    03_roles/roles_catalog.md          — читаемый каталог специальностей

Запуск из корня репозитория:
    python3 tools/build_docs.py
"""
from __future__ import annotations

import csv
import math
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
Q_MD = ROOT / "02_questions" / "questionnaire.md"
D_MD = ROOT / "02_questions" / "documents_request.md"
ROLES_CSV = ROOT / "03_roles" / "roles.csv"
OUT_DOCX = ROOT / "02_questions" / "voprosy_kompanii.docx"
OUT_XLSX = ROOT / "02_questions" / "zayavka_na_podbor.xlsx"
OUT_CATALOG = ROOT / "03_roles" / "roles_catalog.md"

PRIORITY = {"🔴": "1", "🟡": "2", "⚪": "3"}
PRIORITY_FILL = {"1": "F4CCCC", "2": "FFF2CC", "3": "EDEDED"}
PRIORITY_TEXT = {
    "1": "нужно до начала поиска",
    "2": "нужно, чтобы отвечать кандидатам",
    "3": "желательно",
}

# ---------------------------------------------------------------- разбор Markdown

QUESTION_RE = re.compile(r"^- \*\*(\d+\.\d+)\*\*\s*(🔴|🟡|⚪)\s*(.+)$")
WHY_RE = re.compile(r"^\s+_Зачем:_\s*(.+)$")
DOC_RE = re.compile(r"^- \*\*(Д\d+)\*\*\s*(🔴|🟡|⚪)\s*\*\*(.+?)\*\*\s*(.*?)\s*(?:—\s*(.+))?$")


def strip_md(text: str) -> str:
    """Убирает разметку Markdown, оставляя читаемый текст."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)  # ссылки -> текст
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\w)_(.+?)_(?!\w)", r"\1", text)
    return text.strip()


def parse_questionnaire(path: Path) -> dict:
    lines = path.read_text(encoding="utf-8").splitlines()
    data = {"title": "", "meta": [], "howto": [], "short": [], "blocks": [], "appendix": []}
    section = "meta"
    block = None
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("# ") and not data["title"]:
            data["title"] = strip_md(line[2:])
            continue
        if line.startswith("## "):
            head = strip_md(line[3:])
            if head.startswith("Короткий список"):
                section = "short"
                data["short_title"] = head
            elif head.startswith("Блок"):
                section = "block"
                block = {"title": head, "questions": [], "text": []}
                data["blocks"].append(block)
            continue
        if line.strip() == "**Как заполнять**":
            section = "howto"
            continue
        if line.strip() == "**Приложения**":
            section = "appendix"
            continue
        if not line.strip() or line.strip() == "---":
            continue
        if section == "meta" and line.startswith("**"):
            data["meta"].append(strip_md(line))
        elif section == "howto":
            if line.startswith("- "):
                data["howto"].append(strip_md(line[2:]))
            elif data["howto"]:
                data["howto"][-1] += " " + strip_md(line)
        elif section == "short":
            m = re.match(r"^\d+\.\s+(.+)$", line)
            if m:
                data["short"].append(strip_md(m.group(1)))
        elif section == "block" and block is not None:
            m = QUESTION_RE.match(line)
            w = WHY_RE.match(line)
            if m:
                block["questions"].append(
                    {"num": m.group(1), "prio": PRIORITY[m.group(2)], "text": strip_md(m.group(3)), "why": ""}
                )
            elif w and block["questions"]:
                block["questions"][-1]["why"] = strip_md(w.group(1))
            else:
                block["text"].append(strip_md(line))
        elif section == "appendix":
            m = re.match(r"^\d+\.\s+(.+)$", line)
            if m:
                data["appendix"].append(strip_md(m.group(1)))
    return data


def parse_documents(path: Path) -> list[dict]:
    docs = []
    group = ""
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            group = strip_md(line[3:])
            continue
        m = DOC_RE.match(line.strip())
        if m:
            docs.append(
                {
                    "num": m.group(1),
                    "prio": PRIORITY[m.group(2)],
                    "name": strip_md(m.group(3)),
                    "extra": strip_md(m.group(4) or ""),
                    "why": strip_md(m.group(5) or ""),
                    "group": group,
                }
            )
    return docs


def read_roles(path: Path) -> list[dict]:
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


# ---------------------------------------------------------------- Word


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_widths(table, widths_cm: list[float]) -> None:
    """Ширины колонок и в сетке таблицы (LibreOffice, Google Docs), и в ячейках (Word)."""
    table.autofit = False
    for idx, w in enumerate(widths_cm):
        table.columns[idx].width = Cm(w)
    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            row.cells[idx].width = Cm(w)


def cell_text(cell, text: str, bold=False, size=10, align=None, color=None, italic=False):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if align:
        p.alignment = align
    return p


def add_page_number_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("стр. ")
    run.font.size = Pt(8)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2.font.size = Pt(8)
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)


def build_docx(q: dict, docs: list[dict], out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    add_page_number_footer(sec)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    for name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 11)):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    doc.add_paragraph(q["title"], style="Title")
    for m in q["meta"]:
        label, _, value = m.partition(":")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(label + ":")
        r.bold = True
        p.add_run(value)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.add_run("Заполнил(а): ______________________   Должность: __________________   Дата: ____________")

    doc.add_heading("Как заполнять", level=1)
    for item in q["howto"]:
        if item.startswith("Приоритет:"):
            continue
        doc.add_paragraph(item, style="List Bullet")

    legend = doc.add_table(rows=3, cols=2)
    legend.style = "Table Grid"
    for i, key in enumerate(("1", "2", "3")):
        c0, c1 = legend.rows[i].cells
        cell_text(c0, key, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(c0, PRIORITY_FILL[key])
        cell_text(c1, f"приоритет {key} — {PRIORITY_TEXT[key]}")
    set_widths(legend, [0.9, 8.0])

    doc.add_heading(q.get("short_title", "Короткий список"), level=1)
    doc.add_paragraph("Если времени мало — сначала эти вопросы (в скобках номера полных вопросов ниже):")
    for item in q["short"]:
        doc.add_paragraph(item, style="List Number")

    widths = [1.1, 0.8, 8.9, 6.6]
    for block in q["blocks"]:
        doc.add_heading(block["title"], level=1)
        if block["questions"]:
            t = doc.add_table(rows=1, cols=4)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = t.rows[0]
            for c, h in zip(hdr.cells, ("№", "!", "Вопрос", "Ответ")):
                cell_text(c, h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                shade(c, "D9E2F3")
            repeat_header(hdr)
            for qq in block["questions"]:
                row = t.add_row()
                keep_row_together(row)
                c_num, c_pr, c_q, _c_ans = row.cells
                cell_text(c_num, qq["num"], size=9)
                cell_text(c_pr, qq["prio"], bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                shade(c_pr, PRIORITY_FILL[qq["prio"]])
                cell_text(c_q, qq["text"], size=9.5)
                if qq["why"]:
                    wp = c_q.add_paragraph()
                    wp.paragraph_format.space_before = Pt(2)
                    wr = wp.add_run("Зачем: " + qq["why"])
                    wr.italic = True
                    wr.font.size = Pt(8)
                    wr.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            set_widths(t, widths)
        for para in block["text"]:
            p = doc.add_paragraph(para)
            p.runs[0].italic = True

    doc.add_heading("Приложение 1. Таблица-заявка на подбор", level=1)
    doc.add_paragraph(
        "Отдельный файл Excel «zayavka_na_podbor.xlsx»: по каждой нужной специальности укажите количество, "
        "сроки, требования, ставку и условия. Типовые требования уже вписаны — исправьте их под объект."
    )

    doc.add_heading("Приложение 2. Документы, которые нужны для подбора", level=1)
    doc.add_paragraph(
        "Если документа нет, но информация есть — достаточно ответа на вопрос. "
        "Документы с персональными данными сотрудников не нужны."
    )
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0]
    for c, h in zip(hdr.cells, ("№", "!", "Документ", "Для чего нужен", "Отметка о передаче")):
        cell_text(c, h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(c, "D9E2F3")
    repeat_header(hdr)
    for d in docs:
        row = t.add_row()
        keep_row_together(row)
        c = row.cells
        cell_text(c[0], d["num"], size=9)
        cell_text(c[1], d["prio"], bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(c[1], PRIORITY_FILL[d["prio"]])
        p_name = cell_text(c[2], d["name"], bold=True, size=9)
        if d["extra"]:
            sep = "" if d["extra"][:1] in ":;,." else " "
            r_extra = p_name.add_run(sep + d["extra"])
            r_extra.font.size = Pt(8.5)
        cell_text(c[3], d["why"], size=8.5)
    set_widths(t, [1.1, 0.8, 5.6, 7.1, 2.8])

    doc.save(out)


# ---------------------------------------------------------------- Excel

XLSX_COLUMNS = [
    # (заголовок, ширина, источник: поле roles.csv или None)
    ("№", 5, None),
    ("Этап работ", 22, "Этап"),
    ("Специальность", 30, "Специальность"),
    ("Нужна?", 11, None),
    ("Сколько человек", 10, None),
    ("С какой даты", 12, None),
    ("На какой срок, мес.", 10, None),
    ("По одному / бригада / субподряд", 16, None),
    ("Разряд, мин. опыт", 14, None),
    ("Требования и удостоверения (типовые — исправьте под объект)", 42, "Типовые документы и допуски"),
    ("Ставка, ₽", 11, None),
    ("За что ставка", 14, None),
    ("На руки / до НДФЛ", 13, None),
    ("Режим", 11, None),
    ("Часов в смене, график", 14, None),
    ("Оформление", 16, None),
    ("Иностранцы", 16, None),
    ("Срочность", 16, None),
    ("Комментарий", 30, None),
]

VALIDATIONS = {
    "Нужна?": "да,нет,возможно позже",
    "По одному / бригада / субподряд": "по одному,бригада,субподряд (организация),любой вариант",
    "За что ставка": "₽/час,₽/смена,₽/месяц,сдельно (за объём)",
    "На руки / до НДФЛ": "на руки,до вычета НДФЛ",
    "Режим": "вахта,ежедневно,любой",
    "Оформление": "трудовой договор,ГПХ,самозанятый,субподряд ИП/ООО,любой",
    "Иностранцы": "ЕАЭС и патент,только ЕАЭС,нет",
    "Срочность": "1 — срочно,2 — в течение месяца,3 — позже",
}

STAGE_FILLS = ["FFFFFF", "F2F6FC"]


def estimate_lines(text: str, width: float) -> int:
    if not text:
        return 1
    chars_per_line = max(int(width * 1.15), 1)
    return sum(max(1, math.ceil(len(part) / chars_per_line)) for part in str(text).split("\n"))


def build_xlsx(roles: list[dict], out: Path, extra_rows: int = 8) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявка"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(wrap_text=True, vertical="top")

    for col, (title, width, _src) in enumerate(XLSX_COLUMNS, start=1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E2F3")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = border
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 48

    stage_idx, prev_stage = -1, None
    rows = [dict(r) for r in roles] + [
        {"Этап": "Другое", "Специальность": "(впишите)", "Типовые документы и допуски": ""}
        for _ in range(extra_rows)
    ]
    for i, role in enumerate(rows, start=1):
        r = i + 1
        if role["Этап"] != prev_stage:
            stage_idx += 1
            prev_stage = role["Этап"]
        fill = PatternFill("solid", fgColor=STAGE_FILLS[stage_idx % 2])
        max_lines = 1
        for col, (title, width, src) in enumerate(XLSX_COLUMNS, start=1):
            value = i if title == "№" else (role.get(src, "") if src else None)
            c = ws.cell(row=r, column=col, value=value)
            c.alignment = wrap_top
            c.border = border
            c.fill = fill
            if value:
                max_lines = max(max_lines, estimate_lines(str(value), width))
        ws.row_dimensions[r].height = max(15 * max_lines, 18)

    last = len(rows) + 1
    titles = [t for t, _, _ in XLSX_COLUMNS]
    for title, options in VALIDATIONS.items():
        col = get_column_letter(titles.index(title) + 1)
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        dv.error = "Выберите значение из списка или оставьте пустым"
        dv.errorStyle = "information"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(XLSX_COLUMNS))}{last}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    info = wb.create_sheet("Как заполнять")
    info.column_dimensions["A"].width = 110
    lines = [
        ("Таблица-заявка на подбор персонала", True),
        ("", False),
        ("1. Отметьте в колонке «Нужна?» специальности, которые нужны на объекте (выбор из списка).", False),
        ("2. Для нужных специальностей укажите количество, сроки, ставку и условия. Если условия у всех одинаковые "
         "(режим, оформление, иностранцы) — можно заполнить одну строку и написать «как выше».", False),
        ("3. Колонка «Требования и удостоверения» заполнена типовыми требованиями — исправьте под свой объект.", False),
        ("4. Если нужной специальности нет в списке — впишите её в пустые строки внизу («Другое»).", False),
        ("5. Если точного ответа пока нет — напишите диапазон или «уточним к …».", False),
        ("", False),
        ("Общие условия (жильё, питание, доставка, график выплат) — в опроснике «voprosy_kompanii.docx».", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = info.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=13 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out)


# ---------------------------------------------------------------- каталог специальностей


def build_catalog(roles: list[dict], out: Path) -> None:
    stages: dict[str, list[dict]] = {}
    for r in roles:
        stages.setdefault(r["Этап"], []).append(r)

    lines = [
        "# Каталог специальностей для строительства комплекса по переработке отходов",
        "",
        "> Файл собирается автоматически из [`roles.csv`](roles.csv) скриптом `tools/build_docs.py`.",
        "> Правки вносить в `roles.csv`, затем запускать скрипт.",
        "",
        "Это **типовой** перечень: кто может понадобиться на таком объекте по этапам строительства.",
        "Реальная потребность — только после ответа компании (таблица-заявка). Требования к документам —",
        "типовые, их нужно сверять с требованиями компании и заказчика. Рыночные ориентиры — из открытых",
        "объявлений, подробности и источники — в [`reference/market_rates.md`](../reference/market_rates.md).",
        "",
        f"Всего специальностей: **{len(roles)}** по **{len(stages)}** этапам.",
        "",
        "## Оглавление",
        "",
    ]
    for stage, items in stages.items():
        names = ", ".join(i["Специальность"] for i in items)
        lines.append(f"- **{stage}** — {names}")
    lines.append("")

    for stage, items in stages.items():
        lines += [f"## {stage}", ""]
        for r in items:
            kw = ", ".join(f"`{k.strip()}`" for k in r["Запросы для поиска"].split(";") if k.strip())
            lines.append(f"### {r['Специальность']}")
            lines.append("")
            lines.append(f"- **Что делает:** {r['Что делает на объекте']}")
            lines.append(f"- **Документы и допуски (типовые):** {r['Типовые документы и допуски']}")
            lines.append(f"- **Как проверить:** {r['Как проверить']}")
            lines.append(f"- **Запросы для поиска:** {kw}")
            lines.append(f"- **Ориентир по рынку:** {r['Ориентир по рынку (СПб/ЛО)']}")
            if r["Что уточнить у компании"].strip():
                lines.append(f"- **Уточнить у компании:** {r['Что уточнить у компании']}")
            lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    q = parse_questionnaire(Q_MD)
    docs = parse_documents(D_MD)
    roles = read_roles(ROLES_CSV)
    n_q = sum(len(b["questions"]) for b in q["blocks"])
    if n_q == 0 or not docs or not roles:
        raise SystemExit("Ошибка разбора исходников: проверьте формат файлов")
    build_docx(q, docs, OUT_DOCX)
    build_xlsx(roles, OUT_XLSX)
    build_catalog(roles, OUT_CATALOG)
    print(f"Вопросов: {n_q} в {len(q['blocks'])} блоках; коротких: {len(q['short'])}; документов: {len(docs)}; "
          f"специальностей: {len(roles)}")
    for p in (OUT_DOCX, OUT_XLSX, OUT_CATALOG):
        print(f"  -> {p.relative_to(ROOT)} ({p.stat().st_size // 1024} КБ)")


if __name__ == "__main__":
    main()
