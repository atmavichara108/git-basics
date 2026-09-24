#!/usr/bin/env python3
"""Извлечение текста из входящих документов.

По умолчанию обрабатывает все файлы в 00_inbox/ (кроме README.md) и складывает текст
в 01_company/documents/text/<имя>.txt. Страницы PDF без текстового слоя (сканы)
рендерятся в PNG в /tmp/render/<имя>/ — их нужно читать визуально.

Запуск из корня репозитория:
    python3 tools/extract_text.py              # все файлы из 00_inbox
    python3 tools/extract_text.py файл1 файл2  # конкретные файлы
    python3 tools/extract_text.py --move       # после извлечения перенести оригиналы
                                               # в 01_company/documents/original/
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
INBOX = ROOT / "00_inbox"
TEXT_DIR = ROOT / "01_company" / "documents" / "text"
ORIG_DIR = ROOT / "01_company" / "documents" / "original"
RENDER_DIR = Path("/tmp/render")
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic"}
MIN_PAGE_CHARS = 25  # меньше — считаем страницу сканом


def pdf_text(path: Path) -> tuple[str, dict]:
    import pdfplumber

    parts, scanned = [], []
    with pdfplumber.open(path) as pdf:
        n = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            txt = page.extract_text() or ""
            if len(txt.strip()) < MIN_PAGE_CHARS:
                scanned.append(i)
            parts.append(f"=== Страница {i} ===\n{txt}")
            for t_idx, table in enumerate(page.extract_tables() or [], start=1):
                rows = [" | ".join((c or "").replace("\n", " ") for c in row) for row in table]
                parts.append(f"--- Таблица {t_idx} (стр. {i}) ---\n" + "\n".join(rows))
    info = {"страниц": n, "сканов": len(scanned)}
    if scanned:
        info["png"] = render_pages(path, scanned)
    return "\n\n".join(parts), info


def render_pages(path: Path, pages: list[int], scale: float = 2.0) -> str:
    import pypdfium2 as pdfium

    out = RENDER_DIR / path.stem
    out.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(path))
    for p in pages:
        img = pdf[p - 1].render(scale=scale).to_pil()
        img.save(out / f"page-{p:03d}.png")
    return str(out)


def docx_text(path: Path) -> tuple[str, dict]:
    from docx import Document

    d = Document(path)
    parts = []
    for sec in d.sections:
        for p in sec.header.paragraphs:
            if p.text.strip():
                parts.append(f"[колонтитул] {p.text}")
    body = d.element.body
    n_tables = 0
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in child.iter() if t.tag.endswith("}t"))
            if text.strip():
                parts.append(text)
        elif tag == "tbl":
            n_tables += 1
            parts.append(f"--- Таблица {n_tables} ---")
            for tr in child.iter():
                if tr.tag.endswith("}tr"):
                    cells = []
                    for tc in tr.iterchildren():
                        if tc.tag.endswith("}tc"):
                            cells.append(" ".join(
                                (t.text or "") for t in tc.iter() if t.tag.endswith("}t")).strip())
                    parts.append(" | ".join(cells))
    return "\n".join(parts), {"таблиц": n_tables}


def xlsx_text(path: Path) -> tuple[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"=== Лист: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            if any(v not in (None, "") for v in row):
                parts.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(parts), {"листов": len(wb.worksheets)}


def xls_text(path: Path) -> tuple[str, dict]:
    import xlrd

    wb = xlrd.open_workbook(path)
    parts = []
    for sh in wb.sheets():
        parts.append(f"=== Лист: {sh.name} ===")
        for r in range(sh.nrows):
            vals = [str(v) for v in sh.row_values(r)]
            if any(v.strip() for v in vals):
                parts.append("\t".join(vals))
    return "\n".join(parts), {"листов": wb.nsheets}


def doc_text_crude(path: Path) -> tuple[str, dict]:
    """Грубое извлечение из старого .doc: ищем читаемые фрагменты в UTF-16LE и cp1251."""
    raw = path.read_bytes()
    found = []
    for enc in ("utf-16le", "cp1251"):
        text = raw.decode(enc, errors="ignore")
        found += re.findall(r"[А-Яа-яЁёA-Za-z0-9№«»\"'().,:;%\-–— /\n]{20,}", text)
    uniq = list(dict.fromkeys(s.strip() for s in found if s.strip()))
    return "\n".join(uniq), {"внимание": "грубое извлечение — лучше пересохранить в DOCX/PDF"}


def plain_text(path: Path) -> tuple[str, dict]:
    for enc in ("utf-8", "cp1251"):
        try:
            return path.read_text(encoding=enc), {}
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace"), {}


HANDLERS = {
    ".pdf": pdf_text,
    ".docx": docx_text,
    ".xlsx": xlsx_text,
    ".xlsm": xlsx_text,
    ".xls": xls_text,
    ".doc": doc_text_crude,
    ".txt": plain_text,
    ".md": plain_text,
    ".csv": plain_text,
}


def process(path: Path) -> dict:
    ext = path.suffix.lower()
    if ext in IMAGE_EXT:
        return {"файл": path.name, "тип": "изображение", "итог": "читать визуально (read_file)"}
    handler = HANDLERS.get(ext)
    if not handler:
        return {"файл": path.name, "тип": ext, "итог": "формат не поддерживается"}
    text, info = handler(path)
    TEXT_DIR.mkdir(parents=True, exist_ok=True)
    out = TEXT_DIR / f"{path.stem}.txt"
    out.write_text(f"# Источник: {path.name}\n\n{text}", encoding="utf-8")
    shown = out.relative_to(ROOT) if out.resolve().is_relative_to(ROOT) else out
    return {"файл": path.name, "тип": ext, "символов": len(text), **info, "итог": str(shown)}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="*", help="файлы (по умолчанию — всё из 00_inbox)")
    ap.add_argument("--move", action="store_true", help="перенести оригиналы в 01_company/documents/original/")
    args = ap.parse_args()

    files = [Path(f) for f in args.files] or sorted(
        p for p in INBOX.iterdir() if p.is_file() and p.name != "README.md"
    )
    if not files:
        print("Нет файлов для обработки (00_inbox пуст).")
        return
    for f in files:
        try:
            res = process(f)
        except Exception as exc:  # отчёт об ошибке, но продолжаем с остальными
            res = {"файл": f.name, "ошибка": repr(exc)}
        print(" | ".join(f"{k}: {v}" for k, v in res.items()))
        if args.move and "ошибка" not in res and f.parent.resolve() == INBOX.resolve():
            ORIG_DIR.mkdir(parents=True, exist_ok=True)
            shutil.move(str(f), ORIG_DIR / f.name)
    sys.stdout.flush()


if __name__ == "__main__":
    main()
